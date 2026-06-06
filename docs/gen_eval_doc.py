"""Generate docs/SLAM_Perception_Evaluation.docx for Walle GCS.

Grounded in the actual engines:
  - slam_engine.py            (DA V2 depth + ORB VO + log-odds occupancy)
  - reconstruction_engine.py  (Lingbot streaming 3D point cloud)
  - human_tracker_engine.py   (YOLOv8-pose + DeepSort)
  - face_recognition_engine.py(SCRFD + ArcFace, buffalo_l)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)


def h(doc, text, level):
    p = doc.add_heading(text, level=level)
    return p


def body(doc, text, italic=False, color=None, size=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    if size:
        r.font.size = Pt(size)
    return p


def bullet(doc, text):
    doc.add_paragraph(text, style="List Bullet")


def numbered(doc, text):
    doc.add_paragraph(text, style="List Number")


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], "1F3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t


def shade(cell, hex_color):
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


doc = Document()

# Base style
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

# ───────────────────────── Title page ─────────────────────────
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Walle GCS")
r.bold = True
r.font.size = Pt(28)
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("Perception & SLAM Evaluation Plan")
r.font.size = Pt(16)
r.font.color.rgb = GREY

sub2 = doc.add_paragraph()
sub2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub2.add_run("Depth Perception · Object Detection · 2D SLAM · 3D Reconstruction")
r.italic = True
r.font.size = Pt(11)
r.font.color.rgb = GREY

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("\nDocument version 1.0  ·  2026-06-05\n").font.size = Pt(10)
meta.add_run("Ground Control Station — Drone Autonomy Stack").font.size = Pt(10)

doc.add_paragraph()

# ───────────────────────── 1. Purpose ─────────────────────────
h(doc, "1. What We Are Trying To Do", 1)
body(doc,
     "Walle GCS is the ground control station for an autonomous drone. From a single "
     "monocular video feed it must build a usable map of the environment and understand "
     "the people and objects in it, in real time, without stalling the operator GUI. "
     "Four perception subsystems run alongside the video loop:")

table(doc,
      ["Subsystem", "What it produces", "Why it matters"],
      [
        ["Depth perception", "Per-pixel metric depth (DA V2)", "Feeds SLAM unprojection and obstacle ranging"],
        ["2D SLAM", "Log-odds occupancy grid + pose track", "Live navigation / free-space map"],
        ["3D reconstruction", "Fused .ply point cloud", "Survey / inspection deliverable"],
        ["Object detection", "Human tracks + identities", "Situational awareness, follow/track behaviours"],
      ],
      widths=[1.5, 2.4, 2.6])

body(doc,
     "This document defines the evaluation goal for each subsystem: turning "
     "“the map looks about right” into repeatable, numeric pass/fail criteria, "
     "and recording which models we use and how we benchmark them.", italic=True, color=GREY)

# ───────────────────────── 2. Models ─────────────────────────
h(doc, "2. Models Under Evaluation", 1)

h(doc, "2.1 Depth Perception", 2)
body(doc,
     "The SLAM front-end uses Depth Anything V2 (metric variants) selected by scene and "
     "size. The full comparison study (Eval_models/depth_models_comparison.ipynb) also "
     "benchmarked legacy baselines.")
table(doc,
      ["Role", "Model", "Notes"],
      [
        ["Production (indoor)", "Depth-Anything-V2-Metric-Indoor-Small / Base (hf)", "fp16 GPU, input 322px, used in slam_engine.py"],
        ["Production (outdoor)", "Depth-Anything-V2-Metric-Outdoor-Small / Base (hf)", "scene-switched via DEPTH_MODELS_SLAM"],
        ["Baseline (compared)", "Depth Anything V1, DPT-Hybrid, MiDaS", "reference points in the comparison notebook"],
      ],
      widths=[1.7, 3.0, 2.0])

h(doc, "2.2 Object Detection & Identity", 2)
table(doc,
      ["Role", "Model / Method", "Source"],
      [
        ["Human detection + pose", "YOLOv8-pose (yolov8n-pose.pt)", "human_tracker_engine.py"],
        ["Multi-object tracking", "DeepSort (deep_sort_realtime)", "ID-stable tracks across frames"],
        ["Trajectory prediction", "Kalman + motion classifier", "predictive follow behaviour"],
        ["Face detection", "SCRFD-10g (InsightFace buffalo_l, ONNX)", "face_recognition_engine.py"],
        ["Face recognition", "ResNet50 ArcFace (buffalo_l, ONNX)", "cosine match vs. enrolled embeddings"],
      ],
      widths=[1.8, 2.7, 2.2])

h(doc, "2.3 3D Reconstruction", 2)
table(doc,
      ["Role", "Model", "Source"],
      [
        ["Streaming 3D map", "LingBot-Map (lingbot-map-long.pt)", "reconstruction_engine.py → fused .ply"],
      ],
      widths=[1.8, 3.0, 2.0])

# ───────────────────────── 3. Depth eval ─────────────────────────
h(doc, "3. How We Evaluate — Depth Perception", 1)
body(doc,
     "Depth is scored against ground-truth depth (RGB-D capture or LiDAR) using the "
     "standard monocular-depth metrics computed in Eval_models/eval.ipynb. Because the "
     "models are metric, we evaluate without per-image median rescaling.")
table(doc,
      ["Metric", "Meaning", "Target (indoor)"],
      [
        ["AbsRel", "mean |d_pred − d_gt| / d_gt", "≤ 0.10"],
        ["RMSE (m)", "root-mean-square depth error", "≤ 0.40"],
        ["δ₁ (delta_1)", "% pixels with max(d/d*, d*/d) < 1.25", "≥ 0.90"],
        ["δ₂ (d2)", "same, threshold 1.25²", "≥ 0.97"],
        ["Edge accuracy", "depth sharpness at object boundaries", "qualitative, no smearing"],
        ["Latency / FPS", "fp16 GPU inference rate at 322px", "≥ real-time target"],
      ],
      widths=[1.6, 3.2, 1.6])
body(doc,
     "Selection rule: pick the smallest model whose AbsRel / δ₁ clear target at the "
     "required FPS. Small wins unless Base measurably improves SLAM map quality (§5).",
     italic=True, color=GREY)

# ───────────────────────── 4. Detection eval ─────────────────────────
h(doc, "4. How We Evaluate — Object Detection & Identity", 1)

h(doc, "4.1 Detection (YOLOv8-pose)", 2)
table(doc,
      ["Metric", "Meaning", "Target"],
      [
        ["mAP@0.5", "detection accuracy vs. labelled persons", "≥ 0.85 (person class)"],
        ["Precision / Recall", "at conf 0.35 (engine default)", "P ≥ 0.90, R ≥ 0.85"],
        ["Keypoint OKS", "pose-keypoint accuracy", "≥ 0.80"],
        ["Inference FPS", "GPU throughput, cadence-capped", "≥ real-time target"],
      ],
      widths=[1.7, 3.1, 1.6])

h(doc, "4.2 Tracking (DeepSort)", 2)
table(doc,
      ["Metric", "Meaning", "Target"],
      [
        ["MOTA", "multi-object tracking accuracy", "≥ 0.70"],
        ["ID switches", "identity swaps per track over a run", "minimise; report count"],
        ["Track continuity", "% frames a true person holds one ID", "≥ 90%"],
      ],
      widths=[1.7, 3.1, 1.6])

h(doc, "4.3 Face Recognition (SCRFD + ArcFace)", 2)
table(doc,
      ["Metric", "Meaning", "Target"],
      [
        ["Detection recall", "faces found vs. present (multi-scale)", "≥ 0.95"],
        ["Verification TAR@FAR", "true-accept at fixed false-accept rate", "TAR ≥ 0.95 @ FAR 1e-3"],
        ["Rank-1 ID accuracy", "correct identity vs. enrolled gallery", "≥ 0.95"],
        ["Cosine threshold", "operating point on match score", "tuned on held-out set"],
      ],
      widths=[1.8, 3.0, 1.6])

# ───────────────────────── 5. SLAM 2D ─────────────────────────
h(doc, "5. How We Evaluate — 2D SLAM (Occupancy Map)", 1)
body(doc,
     "The 2D map is a GPU log-odds occupancy grid (map_<ts>_logodds.npy). Metrics load the "
     "raw grid (sigmoid → probability; occupied > 0.60, free < 0.40, matching the engine) "
     "and use the per-frame pose telemetry from the stats dict.")
table(doc,
      ["Metric", "How", "Target"],
      [
        ["Scale error", "landmark distance on grid vs. surveyed", "≤ 2%"],
        ["Wall straightness", "RMS residual of fit to a straight wall", "≤ 1 cell"],
        ["Loop-closure gap", "start→end pose drift / path length", "≤ 2%"],
        ["Yaw drift", "heading error after a closed loop", "≤ 5°"],
        ["Map completeness", "swept area marked known", "≥ 90%"],
        ["Free-space precision", "cells marked free that truly are", "≥ 95% (safety)"],
        ["Obstacle recall", "real obstacles marked occupied", "≥ 90% (safety)"],
        ["VO inlier ratio", "Inliers / Matches from stats", "≥ 0.5 sustained"],
      ],
      widths=[1.8, 3.0, 1.6])

# ───────────────────────── 6. SLAM 3D ─────────────────────────
h(doc, "6. How We Evaluate — 3D Reconstruction (Point Cloud)", 1)
body(doc,
     "The fused .ply is registered to a reference scan (ICP) and scored on geometric "
     "fidelity, coverage, and cleanliness.")
table(doc,
      ["Metric", "How", "Target"],
      [
        ["Cloud-to-ref RMS", "nearest-neighbour distance after ICP", "≤ 3 cm indoor"],
        ["Scale error", "known object dimension in cloud vs. truth", "≤ 2%"],
        ["Planarity RMS", "fit plane to flat surface, residual", "≤ 2 cm"],
        ["Completeness", "reference surface within τ of a point", "≥ 85%"],
        ["Outlier rate", "floaters > 5× median NN distance", "≤ 2%"],
        ["Point count / conf", "N points above export confidence cutoff", "report"],
      ],
      widths=[1.8, 3.0, 1.6])

# ───────────────────────── 7. Cross-cutting + verdict ─────────────────────────
h(doc, "7. Cross-Cutting & Verdict", 1)
bullet(doc, "Performance: end-to-end FPS, frame-drop rate (< 10%), CPU/GPU budget, no thermal throttle.")
bullet(doc, "Determinism: run each sequence ×3; key-metric spread < 1%.")
bullet(doc, "Robustness: recovery from hard turns, occlusion, low light — logged.")
body(doc, "Verdict rule:", )
bullet(doc, "PASS — all geometric-accuracy rows green.")
bullet(doc, "FAIL — any single safety row red (free-space precision, obstacle recall, cloud-to-ref).")
bullet(doc, "CONDITIONAL — outputs green but process rows (inliers, FPS) red: works but fragile.")

# ───────────────────────── 8. Test protocol ─────────────────────────
h(doc, "8. Test Protocol", 1)
numbered(doc, "Freeze a test set per scene class: 1 loop, 1 out-and-back, 1 long straight run (indoor + outdoor).")
numbered(doc, "Record sensor, HFOV, scene/variant, lighting, motion, dynamic objects for each run.")
numbered(doc, "Capture ground truth: surveyed distances, reference LiDAR/photogrammetry scan, or loop self-consistency.")
numbered(doc, "Run the pipeline; log the stats dict to CSV per frame; save .npy grid and .ply.")
numbered(doc, "Score with the eval notebooks / scorer; fill the scorecard; attach artifacts.")
numbered(doc, "Track metrics build-over-build for regression trend lines.")

body(doc,
     "Companion: docs/slam_evaluation.md holds the full per-run scorecard template and the "
     "proposed eval_slam.py scorer.", italic=True, color=GREY)

out = "docs/SLAM_Perception_Evaluation.docx"
doc.save(out)
print("wrote", out)
