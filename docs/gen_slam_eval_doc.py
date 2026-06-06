"""Generate docs/SLAM_Map_Evaluation.docx — 2D + 3D SLAM map evaluation only.

Simple, standalone version of sections 5-8 from the perception eval plan.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x55, 0x55, 0x55)


def shade(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(htext)
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade(hdr[i], "1F3A5F")
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Inches(w)
    return t


def body(doc, text, italic=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = italic
    if color:
        r.font.color.rgb = color
    return p


doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(10.5)

# ── Title ──
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("SLAM Map Evaluation")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = NAVY

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sub.add_run("2D Occupancy Grid · 3D Point Cloud")
r.font.size = Pt(13)
r.font.color.rgb = GREY

meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run("Walle GCS  ·  v1.0  ·  2026-06-05").font.size = Pt(10)

doc.add_paragraph()

# ── 1. 2D SLAM ──
doc.add_heading("1. 2D SLAM — Occupancy Map", level=1)
body(doc,
     "The 2D map is a GPU log-odds occupancy grid (map_<ts>_logodds.npy). Metrics load the "
     "raw grid (sigmoid -> probability; occupied > 0.60, free < 0.40, matching the engine) "
     "and use the per-frame pose telemetry from the stats dict.")
table(doc,
      ["Metric", "How", "Target"],
      [
        ["Scale error", "landmark distance on grid vs. surveyed", "<= 2%"],
        ["Wall straightness", "RMS residual of fit to a straight wall", "<= 1 cell"],
        ["Loop-closure gap", "start->end pose drift / path length", "<= 2%"],
        ["Yaw drift", "heading error after a closed loop", "<= 5 deg"],
        ["Map completeness", "swept area marked known", ">= 90%"],
        ["Free-space precision", "cells marked free that truly are", ">= 95% (safety)"],
        ["Obstacle recall", "real obstacles marked occupied", ">= 90% (safety)"],
        ["VO inlier ratio", "Inliers / Matches from stats", ">= 0.5 sustained"],
      ],
      widths=[1.8, 3.2, 1.6])

# ── 2. 3D reconstruction ──
doc.add_heading("2. 3D Reconstruction — Point Cloud", level=1)
body(doc,
     "The fused .ply is registered to a reference scan (ICP) and scored on geometric "
     "fidelity, coverage, and cleanliness.")
table(doc,
      ["Metric", "How", "Target"],
      [
        ["Cloud-to-ref RMS", "nearest-neighbour distance after ICP", "<= 3 cm indoor"],
        ["Scale error", "known object dimension in cloud vs. truth", "<= 2%"],
        ["Planarity RMS", "fit plane to flat surface, residual", "<= 2 cm"],
        ["Completeness", "reference surface within tau of a point", ">= 85%"],
        ["Outlier rate", "floaters > 5x median NN distance", "<= 2%"],
        ["Point count / conf", "N points above export confidence cutoff", "report"],
      ],
      widths=[1.8, 3.2, 1.6])

# ── 3. Cross-cutting & verdict ──
doc.add_heading("3. Cross-Cutting & Verdict", level=1)
doc.add_paragraph("Performance: end-to-end FPS, frame-drop rate (< 10%), CPU/GPU budget, no thermal throttle.", style="List Bullet")
doc.add_paragraph("Determinism: run each sequence x3; key-metric spread < 1%.", style="List Bullet")
doc.add_paragraph("Robustness: recovery from hard turns, occlusion, low light - logged.", style="List Bullet")

p = doc.add_paragraph()
p.add_run("Verdict rule:").bold = True
doc.add_paragraph("PASS - all geometric-accuracy rows green.", style="List Bullet")
doc.add_paragraph("FAIL - any single safety row red (free-space precision, obstacle recall, cloud-to-ref).", style="List Bullet")
doc.add_paragraph("CONDITIONAL - outputs green but process rows (inliers, FPS) red: works but fragile.", style="List Bullet")

# ── 4. Test protocol ──
doc.add_heading("4. Test Protocol", level=1)
doc.add_paragraph("Freeze a test set per scene class: 1 loop, 1 out-and-back, 1 long straight run (indoor + outdoor).", style="List Number")
doc.add_paragraph("Record sensor, HFOV, scene/variant, lighting, motion, dynamic objects for each run.", style="List Number")
doc.add_paragraph("Capture ground truth: surveyed distances, reference LiDAR/photogrammetry scan, or loop self-consistency.", style="List Number")
doc.add_paragraph("Run the pipeline; log the stats dict to CSV per frame; save .npy grid and .ply.", style="List Number")
doc.add_paragraph("Score with the eval notebooks / scorer; fill the scorecard; attach artifacts.", style="List Number")
doc.add_paragraph("Track metrics build-over-build for regression trend lines.", style="List Number")

out = "docs/SLAM_Map_Evaluation.docx"
doc.save(out)
print("wrote", out)
